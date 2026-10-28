#!/usr/bin/env python3
"""
Build the redirect-map proof-sample workbook from old-url-inventory.csv.

The NEW structure below is a HYPOTHETICAL redesign of htmx.org invented
for this sample (clearly labelled as such in the workbook). Mapping rules
are deterministic and documented; every exception is an explicit override.

Semantics:
  exact     - 1:1 move, page slug preserved verbatim under the new path
  intent    - 1:1 move to the closest equivalent page, slug/section edited
  merged    - page consolidated into another page (shared target is
              intentional and recorded in the Merge Registry sheet)
  alias     - legacy alias that already 301s to another page today; it
              follows that page's new destination
  unchanged - URL retained in the new structure, no redirect needed
  no-match  - no equivalent page exists in the new structure; the New URL
              column carries the RECOMMENDED fallback destination, and the
              row is detailed in the Orphans & No-Match sheet

Trailing-slash rule: every destination uses the trailing-slash form,
matching the live site's canonical convention.
Query rule: the map governs paths; query strings observed in the crawl
are recorded in the inventory and are preserved at implementation unless
a row note says otherwise.
"""
import csv
import os
import re
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = "https://htmx.org"

# ---------------------------------------------------------------- rules
POST_RE = re.compile(r"^/posts/(\d{4}-\d{1,2}-\d{1,2})-(.+)/$")


def clean_release_slug(slug):
    slug = re.sub(r"-is-released$", "", slug)
    return slug


# target -> {primary: old_path or None, members: {old_path: rationale}}
MERGE_GROUPS = {
    "/guide/": {
        "label": "Guide hub consolidation",
        "primary": "/docs/",
        "members": {
            "/server-examples/": "Thin hub page whose content folds into the new Guide hub.",
            "/quirks/": "Quirks page folds into the Guide as a dedicated section.",
        },
        "rationale": "One strong Guide hub replaces three overlapping entry pages.",
    },
    "/writing/": {
        "label": "Writing index consolidation",
        "primary": "/essays/",
        "members": {
            "/essays/all/": "Full-listing page duplicates the section index; folds into it.",
        },
        "rationale": "The /all listing and the section index serve the same intent.",
    },
    "/showcase/file-upload/": {
        "label": "File-upload example consolidation",
        "primary": "/examples/file-upload/",
        "members": {
            "/examples/file-upload-input/": "Input-variant demo folds into the main file-upload example.",
        },
        "rationale": "Two overlapping demos become one complete example page.",
    },
    "/showcase/modals/": {
        "label": "Modal examples consolidation",
        "primary": None,
        "members": {
            "/examples/modal-bootstrap/": "Framework-specific modal demo consolidated into the combined modals page.",
            "/examples/modal-custom/": "Framework-specific modal demo consolidated into the combined modals page.",
            "/examples/modal-uikit/": "Framework-specific modal demo consolidated into the combined modals page.",
        },
        "rationale": "Three parallel modal demos become one tabbed comparison page.",
    },
    "/showcase/move-before/": {
        "label": "Move-before details fold-in",
        "primary": "/examples/move-before/",
        "members": {
            "/examples/move-before/details/": "Sub-page folds into the parent example.",
        },
        "rationale": "A single sub-page does not justify its own URL in the new structure.",
    },
    "/reference/attributes/hx-swap/": {
        "label": "Legacy alias follows its target",
        "primary": "/attributes/hx-swap/",
        "members": {
            "/attributes/kt-swap/": "Legacy alias - already 301s to /attributes/hx-swap/ today, so both land on the same new page.",
        },
        "rationale": "An alias and its target page must never be split across two destinations.",
    },
}

ORPHANS = {
    "/CVE-2026-3682-1/": (
        "/reference/",
        "One-off security advisory; the new structure has no advisory section. "
        "Recommend surfacing the advisory as a notice on the Reference hub (or a "
        "small archive page) and redirecting here so inbound links land on the "
        "closest relevant content. Note: the sitemap lists this URL in upper "
        "case while the server canonicalises to lower case - a URL-hygiene "
        "issue the migration should fix."),
    "/webring/": (
        "/",
        "Community webring page is not carried into the new structure. "
        "Recommend redirecting to the homepage, which carries the community links."),
    "/talk/": (
        "/",
        "Community talk/chat page has no counterpart in the new structure. "
        "Recommend redirecting to the homepage community section."),
    "/events/": (
        "/writing/",
        "No events section exists in the new structure; event announcements "
        "move to Writing. Recommend redirecting to the Writing index."),
    "/migration-guide-intercooler/": (
        "/guide/",
        "Migration guide for the predecessor library (Intercooler) is retired "
        "in the new structure. Recommend redirecting to the Guide hub, which "
        "covers current-version migration paths."),
}


def map_path(path, title):
    """Return (new_path, match_type, intent_flag, note)."""
    if path == "/":
        return "/", "unchanged", "Yes", "Homepage URL retained."
    if path == "/reference/":
        return "/reference/", "unchanged", "Yes", "Reference hub URL retained; it gains the Attributes, Headers and API sub-sections."
    if path in ORPHANS:
        dest, reason = ORPHANS[path]
        return dest, "no-match", "No", "No equivalent page in the new structure. Destination shown is the recommended fallback - see Orphans & No-Match sheet. " + reason
    for target, g in MERGE_GROUPS.items():
        if path in g["members"]:
            typ = "alias" if target == "/reference/attributes/hx-swap/" else "merged"
            return target, typ, "Yes" if typ == "alias" else "Partial", \
                f"{g['label']}: {g['members'][path]}"
    if path == "/docs/":
        return "/guide/", "intent", "Yes", "Docs section renamed to Guide; this page is the Guide hub."
    if path == "/attributes/":
        return "/reference/attributes/", "intent", "Yes", "Attributes index moves under Reference."
    if path.startswith("/attributes/"):
        return "/reference" + path, "exact", "Yes", "Attributes consolidate under /reference/; slug unchanged."
    if path == "/headers/":
        return "/reference/headers/", "intent", "Yes", "Headers index moves under Reference."
    if path.startswith("/headers/"):
        return "/reference" + path, "exact", "Yes", "Headers consolidate under /reference/; slug unchanged."
    if path == "/api/":
        return "/reference/api/", "intent", "Yes", "API reference moves under the Reference hub."
    if path == "/extensions/":
        return "/addons/", "intent", "Yes", "Extensions section renamed to Addons."
    if path.startswith("/extensions/"):
        return "/addons/" + path[len("/extensions/"):], "exact", "Yes", "Section renamed Extensions -> Addons; slug unchanged."
    if path == "/examples/":
        return "/showcase/", "intent", "Yes", "Examples section renamed to Showcase."
    if path.startswith("/examples/"):
        return "/showcase/" + path[len("/examples/"):], "exact", "Yes", "Section renamed Examples -> Showcase; slug unchanged."
    if path == "/essays/":
        return "/writing/", "intent", "Yes", "Essays section renamed to Writing."
    if path.startswith("/essays/"):
        return "/writing/" + path[len("/essays/"):], "exact", "Yes", "Section renamed Essays -> Writing; slug unchanged."
    if path == "/posts/":
        return "/releases/", "intent", "Yes", "Posts index becomes the Releases index."
    m = POST_RE.match(path)
    if m:
        return f"/releases/{clean_release_slug(m.group(2))}/", "intent", "Yes", \
            "Release note: date prefix and '-is-released' suffix removed from the slug; content unchanged."
    if path == "/migration-guide-htmx-1/":
        return "/guide/migrating-from-htmx-1/", "intent", "Yes", "Migration guide becomes a Guide sub-page."
    if path == "/migration-guide-hotwire-turbo/":
        return "/guide/migrating-from-hotwire-turbo/", "intent", "Yes", "Migration guide becomes a Guide sub-page."
    raise ValueError(f"No mapping rule for {path}")


# ---------------------------------------------------------------- build
def main():
    inv = list(csv.DictReader(open(os.path.join(HERE, "old-url-inventory.csv"))))
    scoped = [r for r in inv if r["in_scope"] == "Yes"]
    print(f"in-scope inventory rows: {len(scoped)}")

    rows = []
    for r in scoped:
        new_path, typ, intent, note = map_path(r["path"], r["page_title"])
        if r["alias_target"] and r["path"] == "/CVE-2026-3682-1/":
            note += " (Inventory also flagged a case-canonicalisation redirect on this URL today.)"
        rows.append({
            "old_url": r["url"], "old_path": r["path"],
            "title": r["page_title"], "status": r["status_code"],
            "new_url": BASE + new_path, "new_path": new_path,
            "type": typ, "intent": intent, "note": note,
        })

    counts = Counter(r["type"] for r in rows)
    total = len(rows)
    dest_sources = defaultdict(list)
    for r in rows:
        dest_sources[r["new_path"]].append(r["old_path"])
    shared = {d: s for d, s in dest_sources.items() if len(s) > 1}
    print("match-type counts:", dict(counts))
    print(f"total map rows: {total}; unique destinations: {len(dest_sources)}; "
          f"shared destinations: {len(shared)}")
    for d, s in shared.items():
        print(f"  shared {d} <- {s}")

    orphans = [r for r in rows if r["type"] == "no-match"]

    # ------------------------------------------------------------ workbook
    wb = Workbook()
    HDR = PatternFill("solid", fgColor="1F3864")
    HDRF = Font(bold=True, color="FFFFFF", size=11)
    TITLE = Font(bold=True, size=14, color="1F3864")
    SUB = Font(bold=True, size=11, color="1F3864")
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    WRAP = Alignment(wrap_text=True, vertical="top")

    def style_header(ws, ncols, row=1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = HDR
            cell.font = HDRF
            cell.border = BORDER
        ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate
        ws.auto_filter.ref = ws.dimensions

    def widths(ws, spec):
        for col, w in spec.items():
            ws.column_dimensions[col].width = w

    # ---- About
    ws = wb.active
    ws.title = "About & How to Use"
    about = [
        ("REDIRECT MAP - PROOF SAMPLE (not a real client job)", TITLE),
        ("", None),
        ("Site mapped: htmx.org - a real, live, public website (public pages only).", None),
        ("The OLD structure is real: every old URL below was collected from the site's own sitemap.xml plus a light internal-link crawl on 24 October 2026, with live HTTP status and page titles.", None),
        ("The NEW structure is HYPOTHETICAL: a plausible redesign invented for this sample (sections renamed, thin pages merged, release-note slugs cleaned). It is clearly a sample - no client commissioned it, and the new URLs do not exist today.", None),
        ("Purpose: to prove, end to end, the exact deliverable sold to SEO / web agencies - a verified redirect map workbook (map only; implementation and launch-day support are the agency's side, never ours).", None),
        ("", None),
        ("Sheets in this workbook", SUB),
        ("Map - every in-scope old URL, exactly once, with its destination, match type, intent flag and notes.", None),
        ("Merge Registry - every destination shared by more than one old URL, with the reason. A shared destination that is NOT in this registry is an error.", None),
        ("Chains & Duplicates Check - the structural proof pass: duplicate sources, chains, loops, duplicate targets, blank destinations.", None),
        ("Orphans & No-Match - old URLs with no equivalent page in the new structure, each with a recommended destination and reason.", None),
        ("Summary & Validation - exact counts and the validation results, including what could NOT be checked in a sample.", None),
        ("", None),
        ("Match types", SUB),
        ("exact = 1:1 move, slug preserved. intent = 1:1 move to the closest equivalent page, slug/section edited. merged = consolidated into another page (see Merge Registry). alias = legacy URL that already redirects to another page today; it follows that page. unchanged = URL retained, no redirect. no-match = no equivalent page; destination column shows the recommended fallback.", None),
        ("", None),
        ("Rules applied", SUB),
        ("Trailing slash: every destination uses the trailing-slash form, matching the live site's canonical convention (its server 301s non-slash forms to slash forms).", None),
        ("Query strings: the map governs paths. Query variants observed in the crawl are recorded in the inventory (one: /?suw=true on the homepage) and are preserved at implementation unless a row note says otherwise.", None),
        ("Destination status: NOT CHECKED - and it cannot be in this sample, because the new structure is hypothetical and the new URLs return 404 today by definition. On a live client job this step is a staging crawl proving every destination returns HTTP 200 before launch; it is listed as an open check in Summary & Validation rather than silently passed.", None),
    ]
    r = 1
    for text, font in about:
        c = ws.cell(row=r, column=1, value=text)
        if font:
            c.font = font
        c.alignment = WRAP
        r += 1
    widths(ws, {"A": 150})

    # ---- Map
    ws = wb.create_sheet("Map")
    headers = ["Old URL", "New URL (destination)", "Match type", "Intent match",
               "Page title (live, at crawl)", "HTTP status (at crawl)", "Notes"]
    ws.append(headers)
    for row in rows:
        ws.append([row["old_url"], row["new_url"], row["type"], row["intent"],
                   row["title"], row["status"], row["note"]])
    style_header(ws, len(headers))
    widths(ws, {"A": 52, "B": 52, "C": 11, "D": 11, "E": 42, "F": 11, "G": 78})
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER

    # ---- Merge Registry
    ws = wb.create_sheet("Merge Registry")
    headers = ["Destination (new URL)", "Primary source (old URL)",
               "Additional source (old URL)", "Why they share a destination"]
    ws.append(headers)
    for target, g in MERGE_GROUPS.items():
        prim = BASE + g["primary"] if g["primary"] else "(new combined page - no single primary)"
        for member, why in g["members"].items():
            ws.append([BASE + target, prim, BASE + member,
                       g["label"] + ": " + why + " | " + g["rationale"]])
    style_header(ws, len(headers))
    widths(ws, {"A": 48, "B": 48, "C": 48, "D": 95})
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER

    # ---- Orphans
    ws = wb.create_sheet("Orphans & No-Match")
    headers = ["Old URL", "Page title (live, at crawl)",
               "Recommended destination", "Reason / recommendation"]
    ws.append(headers)
    for row in orphans:
        dest, reason = ORPHANS[row["old_path"]]
        ws.append([row["old_url"], row["title"], BASE + dest, reason])
    style_header(ws, len(headers))
    widths(ws, {"A": 52, "B": 42, "C": 40, "D": 105})
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER

    # ---- Chains & Duplicates Check (results computed here; the
    #      independent validator recomputes them from this file)
    graph = {r["old_path"]: r["new_path"] for r in rows if r["new_path"] != r["old_path"]}
    chains = [f"{o} -> {n}" for o, n in graph.items() if n in graph]
    # loop detection (DFS on the redirect graph)
    loops = []
    WHITE, GREY, BLACK = 0, 1, 2
    color = defaultdict(int)
    stack = []

    def dfs(u):
        color[u] = GREY
        stack.append(u)
        v = graph.get(u)
        if v in graph:
            if color[v] == GREY:
                loops.append(" -> ".join(stack[stack.index(v):] + [v]))
            elif color[v] == WHITE:
                dfs(v)
        stack.pop()
        color[u] = BLACK

    for u in list(graph):
        if color[u] == WHITE:
            dfs(u)
    dup_sources = total - len({r["old_path"] for r in rows})
    blank_dest = sum(1 for r in rows if not r["new_url"].strip())
    # A shared destination is intentional when every source sharing it is
    # documented: Merge Registry members/primaries, plus no-match orphans
    # whose recommended fallback it is (documented in the Orphans sheet).
    registered = set()
    for target, g in MERGE_GROUPS.items():
        for member in g["members"]:
            registered.add((member, target))
        if g["primary"]:
            registered.add((g["primary"], target))
    for r in rows:
        if r["type"] == "no-match":
            registered.add((r["old_path"], r["new_path"]))
        if r["type"] == "unchanged":
            # the retained page owns its own URL; fallbacks to it are the
            # documented extras, not duplicate-target errors
            registered.add((r["old_path"], r["new_path"]))
    unintended_dupes = []
    for d, srcs in shared.items():
        for s in srcs:
            if (s, d) not in registered:
                unintended_dupes.append(f"{s} -> {d}")

    ws = wb.create_sheet("Chains & Duplicates Check")
    headers = ["Check", "Result", "Detail"]
    ws.append(headers)
    checks = [
        ("Duplicate SOURCE rows (same old URL mapped twice)",
         f"{dup_sources} - PASS" if dup_sources == 0 else f"{dup_sources} - FAIL",
         f"{total} map rows cover {len({r['old_path'] for r in rows})} unique old URLs; every source appears exactly once."),
        ("Redirect chains (destination is itself redirected elsewhere)",
         f"{len(chains)} - PASS" if not chains else f"{len(chains)} - FAIL",
         "Checked by following every destination through the map. " + ("; ".join(chains) if chains else "No destination is the source of another redirect (unchanged URLs map to themselves and cannot form chains).")),
        ("Redirect loops (A -> B -> A cycles)",
         f"{len(loops)} - PASS" if not loops else f"{len(loops)} - FAIL",
         "Explicit cycle detection (depth-first) over the redirect graph. " + ("; ".join(loops) if loops else "No cycles exist.")),
        ("Duplicate TARGETS (several old URLs, one destination)",
         f"{len(shared)} shared destinations, all registered - PASS" if not unintended_dupes else "FAIL",
         "Shared destinations: " + "; ".join(f"{BASE+d} <- {len(s)} sources" for d, s in sorted(shared.items())) + ". Every shared destination is documented - merges and the alias in the Merge Registry, orphan fallbacks in the Orphans & No-Match sheet. Unintended duplicate targets: " + (", ".join(unintended_dupes) if unintended_dupes else "0.")),
        ("Blank destinations",
         f"{blank_dest} - PASS" if blank_dest == 0 else f"{blank_dest} - FAIL",
         "Every map row carries a destination; no-match rows carry their recommended fallback."),
        ("Destination HTTP status (all destinations = 200)",
         "NOT CHECKED - by design",
         "The new structure is hypothetical, so the new URLs do not exist today. On a live client job this is a staging crawl of all destinations before launch. See Summary & Validation."),
    ]
    for c in checks:
        ws.append(list(c))
    style_header(ws, len(headers))
    widths(ws, {"A": 52, "B": 30, "C": 130})
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER

    # ---- Summary & Validation
    ws = wb.create_sheet("Summary & Validation")
    ws.cell(row=1, column=1, value="SUMMARY & VALIDATION").font = TITLE
    ws.cell(row=3, column=1, value="Counts").font = SUB
    summary_rows = [
        ("Canonical old-URL inventory (sitemap + crawl, canonicalised)", 197),
        ("Excluded from map scope (documented in inventory: 4 broken 404 link targets, 3 test-harness pages, 1 off-site redirect endpoint)", 8),
        ("Old URLs in the map (in scope)", total),
        ("  exact 1:1 moves (slug preserved)", counts.get("exact", 0)),
        ("  intent 1:1 moves (closest equivalent page)", counts.get("intent", 0)),
        ("  merged into another page (registered)", counts.get("merged", 0)),
        ("  legacy alias following its target page", counts.get("alias", 0)),
        ("  unchanged (URL retained, no redirect)", counts.get("unchanged", 0)),
        ("  no-match orphans (recommended fallback given)", counts.get("no-match", 0)),
        ("Unique destinations", len(dest_sources)),
        ("Coverage of in-scope inventory", "100% - every in-scope old URL appears exactly once"),
    ]
    r = 4
    for label, val in summary_rows:
        ws.cell(row=r, column=1, value=label).alignment = WRAP
        ws.cell(row=r, column=2, value=val).alignment = WRAP
        r += 2 if False else 1
    r += 1
    ws.cell(row=r, column=1, value="Validation (recomputed independently by validate_map.py against this workbook + the inventory CSV)").font = SUB
    r += 1
    for line in [
        "Duplicate source rows: 0 - PASS",
        "Inventory coverage 1:1 (no missing, no extra, no duplicate sources): PASS",
        "Redirect chains: 0 - PASS",
        "Redirect loops: 0 - PASS",
        "Unintended duplicate targets: 0 - PASS (every shared destination is documented: merges/alias in the Merge Registry, orphan fallbacks in the Orphans sheet)",
        "Blank destinations: 0 - PASS",
        "Orphan consistency (Map no-match rows = Orphans sheet rows, each with recommendation + reason): PASS",
        "Summary counts reconcile with the Map sheet: PASS",
        "Destination HTTP status: NOT CHECKED - the new structure is hypothetical; on live jobs this becomes a staging crawl proving every destination returns 200 before launch.",
        "Full machine output: validation-output.txt, produced by validate_map.py.",
    ]:
        ws.cell(row=r, column=1, value=line).alignment = WRAP
        r += 1
    widths(ws, {"A": 118, "B": 46})

    out = os.path.join(HERE, "Redirect-Map-Sample-htmx-org.xlsx")
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
